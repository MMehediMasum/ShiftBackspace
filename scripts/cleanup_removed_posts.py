#!/usr/bin/env python3
"""Remove duplicate/off-niche posts per pet_keyword_clusters_v5.xlsx.

Run three times, changing MODE each time:
  1) MODE = "duplicates"        -> redirects + delete duplicate posts
  2) MODE = "noindex-offniche"  -> add noindex to off-niche posts (deploy, wait)
  3) MODE = "delete-offniche"   -> delete off-niche posts after they deindex
"""
import re
from pathlib import Path
from urllib.parse import urlparse

import openpyxl

XLSX_PATH = "pet_keyword_clusters_v5.xlsx"
BLOG_DIR = Path("src/content/blog")
REDIRECTS_PATH = Path("public/_redirects")

MODE = "noindex-offniche"  # "duplicates" | "noindex-offniche" | "delete-offniche"


def slug_from_url(url):
    return urlparse(str(url)).path.strip("/").split("/")[-1]


def norm(text):
    return re.sub(r"\s+", " ", str(text or "").strip().lower())


def build_keyword_to_kept_url(sheet):
    mapping = {}
    for row in sheet.iter_rows(min_row=2, values_only=True):
        _, _pet, _cluster, _topic_type, keyword, post_url = row
        if keyword and post_url:
            mapping[norm(keyword)] = post_url
    return mapping


def main():
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    removed_sheet = wb["Removed Keywords"]
    kept_sheet = wb["All Keywords by Cluster"]
    keyword_to_kept_url = build_keyword_to_kept_url(kept_sheet)

    redirect_lines, skipped = [], []
    deleted = noindexed = 0

    for row in removed_sheet.iter_rows(min_row=2, values_only=True):
        _, _keyword, post_url, reason, kept_instead = row
        if not post_url:
            continue

        slug = slug_from_url(post_url)
        md_path = BLOG_DIR / f"{slug}.md"
        is_duplicate = bool(reason) and reason.lower().startswith(
            ("duplicate", "near-duplicate")
        )

        if MODE == "duplicates" and is_duplicate:
            kept_url = keyword_to_kept_url.get(norm(kept_instead))
            if not kept_url:
                skipped.append({"slug": slug, "why": "kept post not found"})
                continue
            kept_slug = slug_from_url(kept_url)
            redirect_lines.append(f"/{slug}/ /{kept_slug}/ 301")
            if md_path.exists():
                md_path.unlink()
                deleted += 1

        elif MODE == "noindex-offniche" and not is_duplicate:
            if not md_path.exists():
                skipped.append({"slug": slug, "why": "file missing"})
                continue
            text = md_path.read_text(encoding="utf-8")
            if "noindex:" not in text:
                text = re.sub(r"^---\n", "---\nnoindex: true\n", text, count=1)
                md_path.write_text(text, encoding="utf-8")
                noindexed += 1

        elif MODE == "delete-offniche" and not is_duplicate:
            if md_path.exists():
                md_path.unlink()
                deleted += 1

    if redirect_lines:
        REDIRECTS_PATH.parent.mkdir(parents=True, exist_ok=True)
        existing = (
            REDIRECTS_PATH.read_text(encoding="utf-8")
            if REDIRECTS_PATH.exists()
            else ""
        )
        REDIRECTS_PATH.write_text(
            existing.rstrip("\n") + "\n" + "\n".join(redirect_lines) + "\n",
            encoding="utf-8",
        )

    print(f"Mode: {MODE}")
    print(f"Deleted: {deleted}")
    print(f"Noindexed: {noindexed}")
    print(f"Redirect rules added: {len(redirect_lines)}")
    print(f"Skipped: {len(skipped)}")
    for s in skipped[:20]:
        print(" -", s)


if __name__ == "__main__":
    main()
